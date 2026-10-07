import os
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from zoneinfo import ZoneInfo
import pandas as pd
import requests
import yfinance as yf

warnings.filterwarnings("ignore")

# ============================================================
# ⚙️ TELEGRAM AYARLARI (GitHub Secrets'tan okunur)
# ============================================================
TELEGRAM_BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
TELEGRAM_CHAT_ID   = os.environ["TELEGRAM_CHAT_ID"]
# ============================================================

telegram_session = requests.Session()


# ---------------------- TELEGRAM ----------------------
def send_telegram_message(text, parse_mode="HTML"):
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    payload = {
        "chat_id": TELEGRAM_CHAT_ID,
        "text": text,
        "parse_mode": parse_mode,
        "disable_web_page_preview": True,
    }
    try:
        r = telegram_session.post(url, data=payload, timeout=30)
        if not r.ok:
            print(f"⚠️ Telegram yanıtı: {r.status_code} - {r.text[:200]}")
        return r.ok
    except Exception as e:
        print(f"⚠️ Telegram mesaj hatası: {e}")
        return False


def send_telegram_document(file_path, caption=""):
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
    try:
        with open(file_path, "rb") as f:
            files = {"document": f}
            data = {"chat_id": TELEGRAM_CHAT_ID, "caption": caption, "parse_mode": "HTML"}
            r = telegram_session.post(url, files=files, data=data, timeout=120)
        return r.ok
    except Exception as e:
        print(f"⚠️ Telegram dosya hatası: {e}")
        return False


# ---------------------- HİSSE LİSTESİ ----------------------
BIST_SYMBOLS = [
    "A1CAP", "A1YEN", "AAGYO", "ACSEL", "ADEL", "ADESE", "ADGYO", "AEFES", "AFYON", "AGESA",
    "AGHOL", "AGROT", "AGYO", "AHGAZ", "AHSGY", "AKBNK", "AKCNS", "AKENR", "AKFGY", "AKFIS",
    "AKFYE", "AKGRT", "AKHAN", "AKMGY", "AKSA", "AKSEN", "AKSGY", "AKSUE", "AKYHO", "ALARK",
    "ALBRK", "ALCAR", "ALCTL", "ALFAS", "ALGYO", "ALKA", "ALKIM", "ALKLC", "ALTNY",
    "ALVES", "ANELE", "ANGEN", "ANHYT", "ANSGR", "ARASE", "ARCLK", "ARDYZ", "ARENA", "ARFYE",
    "ARMGD", "ARSAN", "ARTMS", "ARZUM", "ASELS", "ASGYO", "ASTOR", "ASUZU", "ATAGY", "ATAKP",
    "ATATP", "ATATR", "ATEKS", "ATLAS", "ATSYH", "AVGYO", "AVHOL", "AVOD", "AVPGY", "AVTUR",
    "AYCES", "AYDEM", "AYEN", "AYES", "AYGAZ", "AZTEK", "BAGFS", "BAHKM", "BAKAB", "BALAT",
    "BALSU", "BANVT", "BARMA", "BASCM", "BASGZ", "BAYRK", "BEGYO", "BERA", "BESLR", "BESTE",
    "BEYAZ", "BFREN", "BIENY", "BIGCH", "BIGEN", "BIGTK", "BIMAS", "BINBN", "BINHO", "BIOEN",
    "BIZIM", "BJKAS", "BLCYT", "BLUME", "BMSCH", "BMSTL", "BNTAS", "BOBET", "BORLS", "BORSK",
    "BOSSA", "BRISA", "BRKO", "BRKSN", "BRKVY", "BRLSM", "BRMEN", "BRSAN", "BRYAT", "BSOKE",
    "BTCIM", "BUCIM", "BULGS", "BURCE", "BURVA", "BVSAN", "BYDNR", "CANTE", "CASA", "CATES",
    "CCOLA", "CELHA", "CEMAS", "CEMTS", "CEMZY", "CEOEM", "CGCAM", "CIMSA", "CLEBI", "CMBTN",
    "CMENT", "CONSE", "COSMO", "CRDFA", "CRFSA", "CUSAN", "CVKMD", "CWENE", "DAGI", "DAPGM",
    "DARDL", "DCTTR", "DENGE", "DERHL", "DERIM", "DESA", "DESPC", "DEVA", "DGATE", "DGGYO",
    "DGNMO", "DIRIT", "DITAS", "DMRGD", "DMSAS", "DNISI", "DOAS", "DOCO", "DOFER",
    "DOFRB", "DOGUB", "DOHOL", "DOKTA", "DSTKF", "DUNYH", "DURDO", "DURKN", "DYOBY", "DZGYO",
    "EBEBK", "ECILC", "ECOGR", "ECZYT", "EDATA", "EDIP", "EFOR", "EGEEN", "EGEGY", "EGEPO",
    "EGGUB", "EGPRO", "EGSER", "EKGYO", "EKIZ", "EKOS", "EKSUN", "ELITE", "EMKEL", "EMNIS",
    "EMPAE", "ENDAE", "ENERY", "ENJSA", "ENKAI", "ENPRA", "ENSRI", "ENTRA", "EPLAS", "ERBOS",
    "ERCB", "EREGL", "ERSU", "ESCAR", "ESCOM", "ESEN", "ETILR", "ETYAT", "EUHOL", "EUKYO",
    "EUPWR", "EUREN", "EUYO", "EYGYO", "FADE", "FENER", "FLAP", "FMIZP", "FONET", "FORMT",
    "FORTE", "FRIGO", "FRMPL", "FROTO", "FZLGY", "GARAN", "GARFA", "GATEG", "GEDIK", "GEDZA",
    "GENIL", "GENKM", "GENTS", "GEREL", "GESAN", "GIPTA", "GLBMD", "GLCVY", "GLRMK", "GLRYH",
    "GLYHO", "GMTAS", "GOKNR", "GOLTS", "GOODY", "GOZDE", "GRNYO", "GRSEL", "GRTHO", "GSDDE",
    "GSDHO", "GSRAY", "GUBRF", "GUNDG", "GWIND", "GZNMI", "HALKB", "HATEK", "HATSN", "HDFGS",
    "HEDEF", "HEKTS", "HKTM", "HLGYO", "HOROZ", "HRKET", "HTTBT", "HUBVC", "HUNER", "HURGZ",
    "ICBCT", "ICUGS", "IDGYO", "IEYHO", "IHAAS", "IHEVA", "IHGZT", "IHLAS", "IHLGM", "IHYAY",
    "IMASM", "INDES", "INFO", "INGRM", "INTEK", "INTEM", "INVEO", "INVES", "ISATR", "ISBIR",
    "ISBTR", "ISCTR", "ISDMR", "ISFIN", "ISGSY", "ISGYO", "ISKPL", "ISKUR", "ISMEN", "ISSEN",
    "ISYAT", "IZENR", "IZFAS", "IZINV", "IZMDC", "JANTS", "KAPLM", "KAREL", "KARSN", "KARTN",
    "KATMR", "KAYSE", "KBORU", "KCAER", "KCHOL", "KENT", "KERVN", "KFEIN", "KGYO", "KIMMR",
    "KLGYO", "KLKIM", "KLMSN", "KLNMA", "KLRHO", "KLSER", "KLSYN", "KLYPV", "KMPUR", "KNFRT",
    "KOCMT", "KONKA", "KONTR", "KONYA", "KOPOL", "KORDS", "KOTON", "KRDMA", "KRDMB", "KRDMD",
    "KRGYO", "KRONT", "KRPLS", "KRSTL", "KRTEK", "KRVGD", "KSTUR", "KTLEV", "KTSKR", "KUTPO",
    "KUVVA", "KUYAS", "KZBGY", "KZGYO", "LIDER", "LIDFA", "LILAK", "LINK", "LKMNH", "LMKDC",
    "LOGO", "LRSHO", "LUKSK", "LXGYO", "LYDHO", "LYDYE", "MAALT", "MACKO", "MAGEN", "MAKIM",
    "MAKTK", "MANAS", "MARBL", "MARKA", "MARMR", "MARTI", "MAVI", "MCARD", "MEDTR", "MEGAP",
    "MEGMT", "MEKAG", "MEPET", "MERCN", "MERIT", "MERKO", "METRO", "MEYSU", "MGROS", "MHRGY",
    "MIATK", "MMCAS", "MNDRS", "MNDTR", "MOBTL", "MOGAN", "MOPAS", "MPARK", "MRGYO", "MRSHL",
    "MSGYO", "MTRKS", "MTRYO", "MZHLD", "NATEN", "NETAS", "NETCD", "NIBAS", "NTGAZ", "NTHOL",
    "NUGYO", "NUHCM", "OBAMS", "OBASE", "ODAS", "ODINE", "OFSYM", "ONCSM", "ONRYT", "ORCAY",
    "ORGE", "ORMA", "OSMEN", "OSTIM", "OTKAR", "OTTO", "OYAKC", "OYAYO", "OYLUM", "OYYAT",
    "OZATD", "OZGYO", "OZKGY", "OZRDN", "OZSUB", "OZYSR", "PAGYO", "PAHOL", "PAMEL", "PAPIL",
    "PARSN", "PASEU", "PATEK", "PCILT", "PEKGY", "PENGD", "PENTA", "PETKM", "PETUN", "PGSUS",
    "PINSU", "PKART", "PKENT", "PLTUR", "PNLSN", "PNSUT", "POLHO", "POLTK", "PRDGS", "PRKAB",
    "PRKME", "PRZMA", "PSDTC", "PSGYO", "QNBFK", "QNBTR", "QUAGR", "RALYH", "RAYSG", "REEDR",
    "RGYAS", "RNPOL", "RODRG", "RTALB", "RUBNS", "RUZYE", "RYGYO", "RYSAS", "SAFKR", "SAHOL",
    "SAMAT", "SANEL", "SANFM", "SANKO", "SARKY", "SASA", "SAYAS", "SDTTR", "SEGMN", "SEGYO",
    "SEKFK", "SEKUR", "SELEC", "SELVA", "SERNT", "SEYKM", "SILVR", "SISE", "SKBNK", "SKTAS",
    "SKYLP", "SKYMD", "SMART", "SMRTG", "SMRVA", "SNGYO", "SNICA", "SNPAM", "SODSN", "SOKE",
    "SOKM", "SONME", "SRVGY", "SUMAS", "SUNTK", "SURGY", "SUWEN", "SVGYO", "TABGD", "TARKM",
    "TATEN", "TATGD", "TAVHL", "TBORG", "TCELL", "TCKRC", "TDGYO", "TEHOL", "TEKTU", "TERA",
    "TEZOL", "TGSAS", "THYAO", "TKFEN", "TKNSA", "TLMAN", "TMPOL", "TMSN", "TNZTP", "TOASO",
    "TRALT", "TRCAS", "TRENJ", "TRGYO", "TRHOL", "TRILC", "TRMET", "TSGYO", "TSKB", "TSPOR",
    "TTKOM", "TTRAK", "TUCLK", "TUKAS", "TUPRS", "TUREX", "TURGG", "TURSG", "UCAYM", "UFUK",
    "ULAS", "ULKER", "ULUFA", "ULUSE", "ULUUN", "UMPAS", "UNLU", "USAK", "VAKBN", "VAKFA",
    "VAKFN", "VAKKO", "VANGD", "VBTYZ", "VERTU", "VERUS", "VESBE", "VESTL", "VKFYO", "VKGYO",
    "VKING", "VRGYO", "VSNMD", "YAPRK", "YATAS", "YAYLA", "YBTAS", "YEOTK", "YESIL", "YGGYO",
    "YIGIT", "YKBNK", "YKSLN", "YUNSA", "YYAPI", "YYLGD", "ZEDUR", "ZERGY", "ZGYO",
    "ZOREN", "ZRGYO"
]
BIST_SYMBOLS = sorted(list(set(BIST_SYMBOLS)))


# ---------------------- TARAMA (orijinaliyle aynı) ----------------------
def scan_symbol(symbol):
    try:
        ticker_symbol = f"{symbol}.IS"
        ticker = yf.Ticker(ticker_symbol)
        df = ticker.history(period="3mo", interval="1d")

        if df is not None and len(df) >= 22:
            # ⚠️ Bozuk/eksik barları temizle (özellikle son bar boş olabiliyor)
            df = df.dropna(subset=['Open', 'High', 'Low', 'Close'])
            if len(df) < 22:
                return None
            df = df.copy()
            df['Vol_SMA21'] = df['Volume'].rolling(window=21).mean()
            last_bar = df.iloc[-1]

            open_p = last_bar['Open']
            high_p = last_bar['High']
            low_p = last_bar['Low']
            close_p = last_bar['Close']
            volume = last_bar['Volume']
            vol_sma21 = last_bar['Vol_SMA21']

            cond1 = close_p > open_p
            cond2 = volume > vol_sma21

            body = close_p - open_p
            total_range = high_p - low_p
            cond3 = (total_range > 0) and (body >= 0.75 * total_range)

            if cond1 and cond2 and cond3:
                body_ratio = (body / total_range) * 100
                vol_ratio = (volume / vol_sma21)

                return {
                    "Hisse": symbol,
                    "Kapanış Fiyatı": round(close_p, 2),
                    "Gövde Oranı (%)": round(body_ratio, 1),
                    "Hacim / 21G Ort (Kat)": round(vol_ratio, 2)
                }
    except Exception as e:
         print(f"❌ {symbol}: {type(e).__name__}: {e}")
    return None


def build_telegram_message(signals, scan_time_str, total_symbols):
    if not signals:
        return (f"📊 <b>BIST Taraması Tamamlandı</b>\n"
                f"⏰ {scan_time_str}\n"
                f"🔎 Taranan: {total_symbols} hisse\n\n"
                f"❌ Kriterlere uyan hisse bulunamadı.")
    lines = [
        "📊 <b>BIST Hacimli Yeşil Mum Taraması</b>",
        f"⏰ <i>{scan_time_str}</i>",
        f"🔎 Taranan: {total_symbols} hisse",
        f"🎯 <b>Bulunan: {len(signals)} hisse</b>",
        "", "━━━━━━━━━━━━━━━━━━━━",
    ]
    for s in signals:
        lines.append(
            f"🟢 <b>{s['Hisse']}</b>\n"
            f"   💰 Fiyat: <code>{s['Kapanış Fiyatı']:.2f}</code>\n"
            f"   📏 Gövde: <code>%{s['Gövde Oranı (%)']:.1f}</code>\n"
            f"   📈 Hacim: <code>{s['Hacim / 21G Ort (Kat)']:.2f}x</code>"
        )
    lines.append("━━━━━━━━━━━━━━━━━━━━")
    lines.append("<i>⚠️ Yatırım tavsiyesi değildir.</i>")
    return "\n".join(lines)


# ---------------------- ANA AKIŞ ----------------------
if __name__ == "__main__":
    scan_time = datetime.now(ZoneInfo("Europe/Istanbul"))
    scan_time_str = scan_time.strftime("%d.%m.%Y %H:%M:%S")

    print("🚀 'yfinance' Tabanlı Hacimli Yeşil Mum Taraması Başlatılıyor...")
    print(f"⏰ Tarama Zamanı: {scan_time_str}")
    print(f"📊 Taranacak Hisse Sayısı: {len(BIST_SYMBOLS)}")
    print("=" * 70)

    signals = []
    processed_count = 0

    with ThreadPoolExecutor(max_workers=10) as executor:
        future_to_symbol = {executor.submit(scan_symbol, sym): sym for sym in BIST_SYMBOLS}

        for future in as_completed(future_to_symbol):
            processed_count += 1
            res = future.result()
            print(f"\rİlerleme: %{int((processed_count / len(BIST_SYMBOLS)) * 100)} ({processed_count}/{len(BIST_SYMBOLS)})", end="", flush=True)
            if res:
                signals.append(res)

    print("\n\n" + "=" * 70)
    print(f"🎯 KRİTERLERE UYGUN HİSSELER ({len(signals)} Adet Bulundu):")
    print("=" * 70)
    print(f"{'Hisse':<10} {'Fiyat':>10} {'Gövde Oranı':>15} {'Hacim / 21G Ort':>18}")
    print("-" * 70)

    if signals:
        signals = sorted(signals, key=lambda x: x['Hacim / 21G Ort (Kat)'], reverse=True)
        for s in signals:
            print(f"🟢 {s['Hisse']:<8} {s['Kapanış Fiyatı']:>10.2f} %{s['Gövde Oranı (%)']:>13.1f} {s['Hacim / 21G Ort (Kat)']:>17.2f}x")

        # --- EXCEL KAYIT BLOĞU ---
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        filename = f"BIST_Tarama_{scan_time.strftime('%Y%m%d_%H%M')}.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tarama Sonuçları"

        ws["A1"] = f"Tarama Zamanı: {scan_time_str}"
        ws["A1"].font = Font(name="Consolas", size=8, italic=True, bold=True, color="555555")

        headers = ["Hisse", "Kapanış Fiyatı", "Gövde Oranı (%)", "Hacim / 21G Ort (Kat)"]
        start_row = 3

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = Font(name="Consolas", size=8, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
        )

        for row_idx, data in enumerate(signals, start=start_row + 1):
            row_values = [data["Hisse"], data["Kapanış Fiyatı"], data["Gövde Oranı (%)"], data["Hacim / 21G Ort (Kat)"]]
            for col_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Consolas", size=8)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "right")

                if col_idx == 2: cell.number_format = '#,##0.00'
                elif col_idx == 3: cell.number_format = '0.0"%"'
                elif col_idx == 4: cell.number_format = '0.00"x"'

        max_row = start_row + len(signals)
        ws.auto_filter.ref = f"A{start_row}:D{max_row}"

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col if cell.row >= start_row)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

        wb.save(filename)
        print("\n" + "=" * 70)
        print(f"📁 Sonuçlar Excel'e kaydedildi: {filename}")

        # --- TELEGRAM GÖNDERİM ---
        print("\n📨 Telegram'a gönderiliyor...")
        if send_telegram_message(build_telegram_message(signals, scan_time_str, len(BIST_SYMBOLS))):
            print("✅ Mesaj gönderildi.")
        if send_telegram_document(filename, caption=f"📎 BIST Tarama - {scan_time_str}"):
            print("✅ Excel gönderildi.")
    else:
        print("Kriterlere uyan herhangi bir hisse bulunamadı.")
        send_telegram_message(build_telegram_message(signals, scan_time_str, len(BIST_SYMBOLS)))

    print("=" * 70)
